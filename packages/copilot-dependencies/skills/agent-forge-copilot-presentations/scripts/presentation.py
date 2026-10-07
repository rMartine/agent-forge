"""Create and inspect editable PPTX decks; requires python-pptx."""
import argparse
import json
from pathlib import Path


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    inspect = sub.add_parser("inspect")
    inspect.add_argument("input", type=Path)
    create = sub.add_parser("create")
    create.add_argument("brief", type=Path)
    create.add_argument("output", type=Path)
    args = parser.parse_args()
    from pptx import Presentation
    from pptx.util import Inches
    if args.command == "inspect":
        deck = Presentation(args.input)
        slides = [{"number": i + 1, "shapes": [{"name": s.name, "text": s.text if s.has_text_frame else None, "x": s.left, "y": s.top, "width": s.width, "height": s.height} for s in slide.shapes], "notes": slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else None} for i, slide in enumerate(deck.slides)]
        print(json.dumps({"width": deck.slide_width, "height": deck.slide_height, "slides": slides}, ensure_ascii=False, indent=2))
        return
    if args.output.exists():
        parser.error("Output already exists; choose a new output path")
    brief = json.loads(args.brief.read_text(encoding="utf-8-sig"))
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333333), Inches(7.5)
    deck.core_properties.title = brief.get("title", "")
    for item in brief["slides"]:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = item["title"]
        frame = slide.placeholders[1].text_frame
        frame.clear()
        for index, text in enumerate(item.get("bullets", [])):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = str(text)
        if item.get("notes"):
            slide.notes_slide.notes_text_frame.text = item["notes"]
    args.output.parent.mkdir(parents=True, exist_ok=True)
    deck.save(args.output)
    print(str(args.output.resolve()))


if __name__ == "__main__":
    main()
