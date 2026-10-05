"""Create or inspect XLSX; formulas are preserved but never evaluated."""
import argparse
import json
from pathlib import Path


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    inspect = sub.add_parser("inspect")
    inspect.add_argument("input", type=Path)
    inspect.add_argument("--max-rows", type=int, default=30)
    create = sub.add_parser("create")
    create.add_argument("brief", type=Path)
    create.add_argument("output", type=Path)
    args = parser.parse_args()
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font
    if args.command == "inspect":
        if args.max_rows < 1:
            parser.error("max-rows must be positive")
        book = load_workbook(args.input, read_only=True, data_only=False)
        try:
            result = [{"name": sheet.title, "max_row": sheet.max_row, "max_column": sheet.max_column, "rows": list(sheet.iter_rows(max_row=min(args.max_rows, sheet.max_row), values_only=True))} for sheet in book]
            print(json.dumps(result, default=str, ensure_ascii=False, indent=2))
        finally:
            book.close()
        return
    if args.output.exists():
        parser.error("Output already exists; choose a new output path")
    brief = json.loads(args.brief.read_text(encoding="utf-8-sig"))
    if not brief.get("sheets"):
        parser.error("At least one sheet is required")
    book = Workbook()
    book.remove(book.active)
    for item in brief["sheets"]:
        sheet = book.create_sheet(item["name"])
        for row in item["rows"]:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = min(45, max(12, max(len(str(c.value or "")) for c in col) + 2))
    args.output.parent.mkdir(parents=True, exist_ok=True)
    book.save(args.output)
    print(json.dumps({"path": str(args.output.resolve()), "formulasEvaluated": False}))


if __name__ == "__main__":
    main()
