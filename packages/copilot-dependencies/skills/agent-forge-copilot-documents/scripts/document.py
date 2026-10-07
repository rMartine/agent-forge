"""Create or inspect DOCX; requires python-docx. No network operations."""
import argparse
import json
from pathlib import Path


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    inspect = sub.add_parser("inspect")
    inspect.add_argument("input", type=Path)
    create = sub.add_parser("create")
    create.add_argument("brief", type=Path)
    create.add_argument("output", type=Path)
    args = parser.parse_args()
    from docx import Document
    if args.command == "inspect":
        doc = Document(args.input)
        print(json.dumps({"paragraphs": [{"index": i, "text": p.text, "style": p.style.name if p.style else None} for i, p in enumerate(doc.paragraphs)], "tables": [[[c.text for c in r.cells] for r in t.rows] for t in doc.tables]}, ensure_ascii=False, indent=2))
        return
    if args.output.exists():
        parser.error("Output already exists; choose a new output path")
    brief = json.loads(args.brief.read_text(encoding="utf-8-sig"))
    doc = Document()
    if brief.get("title"):
        doc.add_heading(brief["title"], 0)
    for paragraph in brief.get("paragraphs", []):
        doc.add_paragraph(paragraph["text"], style=paragraph.get("style", "Normal"))
    for rows in brief.get("tables", []):
        if not rows or not rows[0] or any(len(row) != len(rows[0]) for row in rows):
            parser.error("Tables must be nonempty rectangular arrays")
        table = doc.add_table(rows=len(rows), cols=len(rows[0]))
        table.style = "Table Grid"
        for i, row in enumerate(rows):
            for j, value in enumerate(row):
                table.cell(i, j).text = str(value)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    print(str(args.output.resolve()))


if __name__ == "__main__":
    main()
